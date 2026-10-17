"""
WhatsApp lead capture
=====================
Every message a customer sends to your WhatsApp Business number (Meta Cloud API)
arrives at the webhook below. Each phone number becomes one lead; the destination,
dates, people and budget are picked out of the chat and filled in automatically.
Fields you edit by hand in the admin panel are never overwritten.

Public (called by Meta)
    GET  /api/whatsapp/webhook   -> one-time verification handshake
    POST /api/whatsapp/webhook   -> incoming messages (signature checked)

Admin (Bearer token)
    GET  /api/admin/whatsapp-leads               -> all leads, newest activity first
    PUT  /api/admin/whatsapp-leads/<id>          -> edit status / notes / fields
    GET  /api/admin/whatsapp-leads/export.xlsx   -> Excel download
"""
import hashlib
import hmac
import io
import json
import logging
import re
from datetime import datetime, timezone

import requests
from flask import Blueprint, current_app, jsonify, request, send_file

from backend.auth import admin_required
from backend.extensions import db
from backend.models import WhatsappLead, WhatsappMessage

log = logging.getLogger(__name__)
whatsapp_bp = Blueprint("whatsapp", __name__)

GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"
FIELDS = ("name", "destination", "travel_dates", "people", "budget")
STATUSES = ("new", "quoted", "follow-up", "booked", "lost")
PLACES = ["Srinagar", "Pahalgam", "Gulmarg", "Sonamarg", "Doodhpathri", "Yusmarg", "Gurez", "Dal Lake",
          "Kashmir", "Leh", "Ladakh", "Betaab", "Aru", "Kokernag", "Verinag"]


# ── Pulling details out of the chat ───────────────────────────────────────

def _regex_extract(text: str) -> dict:
    out = {}
    found = [p for p in PLACES if re.search(r"\b" + re.escape(p) + r"\b", text, re.I)]
    if found:
        out["destination"] = ", ".join(found)
    m = re.search(r"(\d{1,2}\s*(?:st|nd|rd|th)?\s*(?:to|-|till)\s*\d{1,2}\s*(?:st|nd|rd|th)?\s*"
                  r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*)", text, re.I)
    if m:
        out["travel_dates"] = m.group(1)
    m = re.search(r"(\d{1,2})\s*(?:pax|persons?|people|adults?|members?)", text, re.I)
    if m:
        out["people"] = f"{m.group(1)} people"
    m = re.search(r"(?:budget|rs\.?|inr|₹)\s*[:\-]?\s*(\d[\d,]*\s*(?:k|lakh|lac)?)", text, re.I)
    if m:
        out["budget"] = "₹" + m.group(1).strip()
    return out


def _ai_extract(text: str) -> dict:
    key = current_app.config.get("GROQ_API_KEY")
    if not key:
        return {}
    prompt = (
        "You read a WhatsApp chat between a Kashmir travel agency and a customer. "
        "Return ONLY JSON with keys name, destination, travel_dates, people, budget. "
        "Use the customer's own words, short. Use an empty string when the chat does not say. Never guess.\n\n" + text
    )
    try:
        r = requests.post(
            GROQ_URL,
            headers={"Authorization": f"Bearer {key}"},
            json={"model": current_app.config["ASSISTANT_MODEL"],
                  "messages": [{"role": "user", "content": prompt}],
                  "response_format": {"type": "json_object"}, "temperature": 0},
            timeout=10,
        )
        r.raise_for_status()
        data = json.loads(r.json()["choices"][0]["message"]["content"])
        return {k: str(data.get(k) or "").strip()[:200] for k in FIELDS}
    except Exception as exc:  # extraction is a bonus; the lead is saved either way
        log.warning("whatsapp: AI extraction skipped: %s", exc)
        return {}


def _fill_blanks(lead: WhatsappLead):
    recent = (WhatsappMessage.query.filter_by(lead_id=lead.id)
              .order_by(WhatsappMessage.id.desc()).limit(15).all())
    text = "\n".join(m.text for m in reversed(recent) if m.text)
    found = {**_regex_extract(text), **{k: v for k, v in _ai_extract(text).items() if v}}
    for k, v in found.items():
        if k in FIELDS and v and not getattr(lead, k):     # never overwrite what is already there
            setattr(lead, k, v)


# ── Webhook ───────────────────────────────────────────────────────────────

@whatsapp_bp.get("/whatsapp/webhook")
def verify():
    token = current_app.config["WHATSAPP_VERIFY_TOKEN"]
    if token and request.args.get("hub.mode") == "subscribe" and request.args.get("hub.verify_token") == token:
        return request.args.get("hub.challenge", ""), 200
    return "forbidden", 403


def _signature_ok() -> bool:
    secret = current_app.config["WHATSAPP_APP_SECRET"]
    if not secret:
        return False
    sent = request.headers.get("X-Hub-Signature-256", "").removeprefix("sha256=")
    good = hmac.new(secret.encode(), request.get_data(), hashlib.sha256).hexdigest()
    return hmac.compare_digest(sent, good)


def _message_text(msg: dict) -> str:
    kind = msg.get("type")
    if kind == "text":
        return msg["text"].get("body", "")
    if kind in ("image", "video", "document"):
        return (msg.get(kind, {}).get("caption") or "") or f"[{kind}]"
    if kind == "button":
        return msg["button"].get("text", "")
    if kind == "interactive":
        i = msg["interactive"]
        return (i.get("button_reply") or i.get("list_reply") or {}).get("title", "")
    return f"[{kind}]"


@whatsapp_bp.post("/whatsapp/webhook")
def receive():
    if not _signature_ok():
        return "forbidden", 403
    payload = request.get_json(silent=True) or {}
    touched = {}
    try:
        for entry in payload.get("entry", []):
            for change in entry.get("changes", []):
                value = change.get("value", {})
                names = {c.get("wa_id"): c.get("profile", {}).get("name", "") for c in value.get("contacts", [])}
                for msg in value.get("messages", []):
                    phone = msg.get("from", "")
                    wamid = msg.get("id")
                    if not phone or (wamid and WhatsappMessage.query.filter_by(wamid=wamid).first()):
                        continue
                    ts = datetime.fromtimestamp(int(msg.get("timestamp", 0)) or 0, tz=timezone.utc) \
                        if msg.get("timestamp") else datetime.now(timezone.utc)
                    lead = WhatsappLead.query.filter_by(phone=phone).first()
                    if not lead:
                        lead = WhatsappLead(phone=phone, name=names.get(phone, ""), first_at=ts, msg_count=0)
                        db.session.add(lead)
                        db.session.flush()
                    elif not lead.name and names.get(phone):
                        lead.name = names[phone]
                    text = _message_text(msg)
                    db.session.add(WhatsappMessage(wamid=wamid, lead_id=lead.id, text=text, created_at=ts))
                    lead.msg_count = (lead.msg_count or 0) + 1
                    lead.last_message = text[:500]
                    lead.last_at = ts
                    touched[lead.id] = lead
        db.session.commit()
        for lead in touched.values():
            _fill_blanks(lead)
        db.session.commit()
    except Exception:
        db.session.rollback()
        log.exception("whatsapp: could not store incoming message")
        return "error", 500     # Meta retries, and duplicates are ignored by message id
    return "ok", 200


# ── Admin ─────────────────────────────────────────────────────────────────

def _row(l: WhatsappLead) -> dict:
    return {
        "id": l.id, "phone": l.phone, "name": l.name or "", "destination": l.destination or "",
        "travel_dates": l.travel_dates or "", "people": l.people or "", "budget": l.budget or "",
        "status": l.status or "new", "notes": l.notes or "", "last_message": l.last_message or "",
        "msg_count": l.msg_count or 0,
        "first_at": l.first_at.isoformat() if l.first_at else None,
        "last_at": l.last_at.isoformat() if l.last_at else None,
    }


@whatsapp_bp.get("/admin/whatsapp-leads")
@admin_required
def list_leads():
    leads = WhatsappLead.query.order_by(WhatsappLead.last_at.desc()).all()
    cfg = current_app.config
    return jsonify({
        "leads": [_row(l) for l in leads],
        "connected": bool(cfg["WHATSAPP_VERIFY_TOKEN"] and cfg["WHATSAPP_APP_SECRET"]),
    })


@whatsapp_bp.put("/admin/whatsapp-leads/<int:lead_id>")
@admin_required
def update_lead(lead_id):
    lead = db.session.get(WhatsappLead, lead_id)
    if not lead:
        return jsonify({"error": "Lead not found."}), 404
    data = request.get_json(silent=True) or {}
    if "status" in data:
        if data["status"] not in STATUSES:
            return jsonify({"error": "Unknown status."}), 400
        lead.status = data["status"]
    for k in (*FIELDS, "notes"):
        if k in data:
            setattr(lead, k, str(data[k] or "")[:2000 if k == "notes" else 200])
    db.session.commit()
    return jsonify(_row(lead))


@whatsapp_bp.get("/admin/whatsapp-leads/export.xlsx")
@admin_required
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "WhatsApp Leads"
    head = ["First message", "Last message at", "Name", "Phone", "Destination", "Travel dates",
            "People", "Budget", "Status", "Messages", "Last message", "Notes"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A4F")
    for l in WhatsappLead.query.order_by(WhatsappLead.last_at.desc()).all():
        ws.append([
            l.first_at.strftime("%d-%b-%Y %H:%M") if l.first_at else "",
            l.last_at.strftime("%d-%b-%Y %H:%M") if l.last_at else "",
            l.name, "+" + l.phone, l.destination, l.travel_dates, l.people, l.budget,
            l.status, l.msg_count, l.last_message, l.notes,
        ])
    for col, w in zip("ABCDEFGHIJKL", (18, 18, 20, 16, 24, 20, 12, 14, 12, 10, 50, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="whatsapp-leads.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
